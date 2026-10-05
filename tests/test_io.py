"""Importers, the .knoise document and the KherveFitting export."""

import os

import numpy as np
import pytest

from khervenoise import importers
from khervenoise.document import Document, axis_labels, make_spectrum, x_reversed
from khervenoise.exporters import write_khervefitting_xlsx, write_text


def test_vamas(vamas_path):
    res = importers.read_any(vamas_path)
    assert [s['Name'] for s in res.spectra] == ['Pt4f']
    sp = res.spectra[0]
    assert len(sp['B.E.']) == len(sp['Raw Data']) == 281
    assert sp['B.E.'][0] == pytest.approx(90.0)
    assert max(sp['Raw Data']) > 1e4
    assert sp['ExperimentalInfo']['Species & Transition'].startswith('Pt')


@pytest.mark.parametrize("raw,expected", [
    ("Fe LM2", "Felmm"), ("C KL1", "Ckll"), ("V.B.", "VB"), ("O1s", "O1s")])
def test_auger_valence_names(raw, expected):
    assert importers.normalize_auger_and_valence_names(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    ("C 1s", "C1s"), ("C1s Scan", "C1s"), ("Survey", "Survey"),
    ("XPS Survey 2", "Survey2"), ("Valence", "VB"), ("Fe2p3", "Fe2p3")])
def test_sheet_names(raw, expected):
    assert importers.normalize_sheet_name(raw) == expected
    assert importers.phi_region_base("Fe2p3") == "Fe2p"
    assert importers.phi_region_base("C1s") == "C1s"


def test_asc_and_csv(tmp_path):
    asc = tmp_path / "C1s.txt"
    asc.write_text("# header\nBinding Energy\tCounts\n" +
                   "\n".join(f"{295 - i * 0.1:.2f}\t{1000 + i}" for i in range(50)))
    res = importers.read_any(str(asc))
    assert res.spectra[0]['Name'] == "C1s" and len(res.spectra[0]['B.E.']) == 50
    csv = tmp_path / "O1s.csv"
    csv.write_text("BE,Intensity\n" + "\n".join(f"{535 - i * 0.1},{200 + i}" for i in range(40)))
    res = importers.read_any(str(csv))
    assert res.spectra[0]['Name'] == "O1s" and len(res.spectra[0]['Raw Data']) == 40
    semi = tmp_path / "Ti2p.asc"
    semi.write_text("\n".join(f"{470 - i};{5 + i}" for i in range(10)))
    assert len(importers.read_any(str(semi)).spectra[0]['B.E.']) == 10


def test_excel_roundtrip_through_khervefitting_layout(tmp_path, vamas_path):
    sp = importers.read_vamas(vamas_path).spectra[0]
    den = dict(sp, Name="Pt4f_vmd", Denoise={'source': 'Pt4f', 'params': {'method': 'VMD'}})
    path = str(tmp_path / "out.xlsx")
    assert write_khervefitting_xlsx(path, [sp, den]) == ["Pt4f", "Pt4f_vmd"]
    import openpyxl
    ws = openpyxl.load_workbook(path)["Pt4f"]
    assert [c.value for c in ws[1][:4]] == ["BE", "Corrected Data", "Raw Data", "Transmission"]
    assert ws.cell(row=1, column=50).value == "Experimental Description"
    back = importers.read_any(path)
    assert [s['Name'] for s in back.spectra] == ["Pt4f", "Pt4f_vmd"] and not back.dismissed
    np.testing.assert_allclose(back.spectra[0]['Raw Data'],
                               [float(f"{v:.2f}") for v in sp['Raw Data']])


def test_excel_dismisses_like_khervefitting(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "C1s"
    ws.append(["Binding Energy", "Raw Data"])
    for i in range(20):
        ws.append([290 - i * 0.1, 100 + i])
    bad = wb.create_sheet("Notes"); bad.append(["hello", "world"]); bad.append([1, 2])
    wb.create_sheet("Sheet1").append(["x", "y"])
    wb.create_sheet("Results Table")
    path = str(tmp_path / "kf.xlsx"); wb.save(path)
    res = importers.read_any(path)
    assert [s['Name'] for s in res.spectra] == ["C1s"]
    assert {n for n, _r in res.dismissed} == {"Notes", "Sheet1"}


def test_generic_excel_fallback(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(30):
        ws.append([i, i * i])
    path = str(tmp_path / "raw.xlsx"); wb.save(path)
    res = importers.read_any(path)
    assert [s['Name'] for s in res.spectra] == ["raw"]


def test_document_roundtrip(tmp_path):
    doc = Document()
    assert doc.add(make_spectrum("C1s", [1, 2, 3, 4], [5, 6, 7, 8])) == "C1s"
    assert doc.add(make_spectrum("C1s", [1, 2, 3, 4], [5, 6, 7, 8])) == "C1s1"
    assert doc.next_sheet_name("C1s", "VMD") == "C1s_vmd"
    assert doc.rename("C1s", "Carbon") and "Carbon" in doc.names()
    path = str(tmp_path / "p.knoise")
    doc.save(path)
    back = Document.load(path)
    assert back.names() == doc.names() and not back.dirty
    with pytest.raises(ValueError):
        (tmp_path / "x.knoise").write_text("{}")
        Document.load(str(tmp_path / "x.knoise"))


def test_axis_conventions():
    assert axis_labels("C1s")[0] == "Binding Energy (eV)" and x_reversed("C1s")
    assert axis_labels("EELS~LL")[0] == "Energy Loss (eV)" and not x_reversed("EELS~LL")
    assert x_reversed("FTIR_1")
    assert not x_reversed("Raman1")


def test_write_text(tmp_path):
    sp = make_spectrum("A", [1, 2], [3, 4])
    path = str(tmp_path / "a.csv")
    write_text(path, [sp])
    assert open(path).read().splitlines()[0] == "A X,A Y"
    write_text(path, [make_spectrum("C1s", [1, 2], [3, 4])])
    assert open(path).read().splitlines()[0] == "C1s BE,C1s Intensity (CPS)"


def test_every_vendor_module_registers():
    """A vendor module that fails to import is skipped by the registry at
    run time — so the tests must catch it."""
    import importlib
    from khervenoise import vendors
    for name in vendors.MODULES:
        mod = importlib.import_module(f"khervenoise.vendors.{name}")
        assert mod.FORMATS, name
        for fmt in mod.FORMATS:
            assert fmt.extensions and all(e.startswith('.') and e == e.lower()
                                          for e in fmt.extensions), fmt.key
    keys = [f.key for f in vendors.all_formats()]
    assert len(keys) == len(set(keys))


def test_plain_files_are_not_claimed_by_vendors(tmp_path):
    from khervenoise import importers
    for name, text in (("C1s.txt", "290 1\n289 2\n288 3\n287 4\n"),
                       ("C1s.dat", "290 1\n289 2\n288 3\n287 4\n"),
                       ("C1s.csv", "290,1\n289,2\n288,3\n287,4\n")):
        p = tmp_path / name
        p.write_text(text)
        assert importers.kind_of(str(p)) == "data", name
        assert importers.read_any(str(p)).spectra[0]['Name'] == "C1s"


def test_generic_axes_are_not_xps(tmp_path):
    from khervenoise.importers import read_any
    # Arbitrary data is not XPS: generic or header labels, read low to high.
    assert axis_labels("Data1") == ("X", "Y") and not x_reversed("Data1")
    for name in ("C1s", "O 1s", "Ti2p3/2", "Survey", "C KLL", "VB"):
        assert axis_labels(name)[0] == "Binding Energy (eV)", name
    sp = make_spectrum("run", [1], [2], technique="XPS")
    assert x_reversed("run", sp)
    csv = tmp_path / "trace.csv"
    csv.write_text("Time (s),Voltage (V)\n" +
                   "".join(f"{i},{i * i}\n" for i in range(10)))
    sp = read_any(str(csv)).spectra[0]
    assert axis_labels(sp['Name'], sp) == ("Time (s)", "Voltage (V)")
    assert not x_reversed(sp['Name'], sp)
    txt = tmp_path / "scan.txt"
    txt.write_text("Binding Energy\tCounts\n" +
                   "".join(f"{290 - i * .1}\t{i}\n" for i in range(10)))
    sp = read_any(str(txt)).spectra[0]
    assert axis_labels(sp['Name'], sp)[0] == "Binding Energy (eV)"
    assert x_reversed(sp['Name'], sp)
