"""Thermo formats (Avantage .xlsx/.xls, VGD, AVG) against KherveFitting.

Real sample files are referenced on this machine only (they are private /
not redistributable) and skipped when absent.  Where KherveFitting's source
is present, its own functions are loaded with ``ast`` (wx and
``libraries.*`` stubbed) and the spectra are compared value for value.
"""

import ast
import os
import re
import shutil
import struct
import subprocess
import sys
import types

import numpy as np
import pytest

from khervenoise import vendors
from khervenoise.vendors import avantage, avg, vgd

KF_LIB = os.environ.get("KHERVEFITTING_SRC", "/Users/gkerherv/Documents/PycharmProjects/KherveFittingPro") + "/libraries"
RAW = (os.environ.get("KHERVEFITTING_SRC", "/Users/gkerherv/Documents/PycharmProjects/KherveFittingPro") + "/Data-Examples/"
       "zz RawData_ToBeImported")
AVANTAGE = {
    "Cu_Metal_Avantage.xlsx": ["Survey", "O1s", "Cu2p", "C1s"],
    "PET25_07_Avantage.xlsx": ["Survey", "O1s", "C1s"],
    "Ag_Metal_Avantage.xlsx": ["Survey", "O1s", "C1s", "Ag3d"],
    "Au_Metal_Avantage.xlsx": ["Survey", "O1s", "C1s", "Au4f"],
}
KF_WORKBOOK = os.environ.get("KHERVEFITTING_SRC", "/Users/gkerherv/Documents/PycharmProjects/KherveFittingPro") + "/Data/STO_2doublets.xlsx"
VGD_ROOT = "/Users/gkerherv/Documents/KherveFitting/XPS"

need_kf = pytest.mark.skipif(not os.path.isdir(KF_LIB), reason="KherveFitting source not here")


def _vgd_files():
    if not os.path.isdir(VGD_ROOT):
        return []
    out = subprocess.run(["find", VGD_ROOT, "-iname", "*.vgd"],
                         capture_output=True, text=True).stdout
    return sorted(p for p in out.splitlines() if p)


VGD_FILES = _vgd_files()


# ---------------------------------------------------------------------------
# KherveFitting reference: selected top-level functions, executed as-is
# ---------------------------------------------------------------------------
def kf_functions(relpath, names, **extra):
    import json
    import openpyxl
    path = os.path.join(KF_LIB, relpath)
    with open(path, encoding="utf-8") as fh:
        tree = ast.parse(fh.read())
    body = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in names]
    assert {n.name for n in body} == set(names), "KherveFitting function missing"
    ns = {"os": os, "re": re, "struct": struct, "np": np, "json": json,
          "openpyxl": openpyxl, "__name__": "kf_reference"}
    ns.update(extra)
    exec(compile(ast.Module(body=body, type_ignores=[]), path, "exec"), ns)
    return ns


def _stub_libraries(monkeypatch, **modules):
    """sys.modules entries for 'libraries.X.Y' imports done inside functions."""
    for pkg in ("libraries", "libraries.FileMenu"):
        monkeypatch.setitem(sys.modules, pkg, types.ModuleType(pkg))
    for name, attrs in modules.items():
        mod = types.ModuleType(name)
        mod.__dict__.update(attrs)
        monkeypatch.setitem(sys.modules, name, mod)


def kf_read_back(xlsx, sheet_names):
    """add_core_level_Data for each sheet, then refresh_sheets' renaming and
    its reload of 'B.E.' / 'Raw Data' from columns A / B."""
    import pandas as pd
    cfg = kf_functions("ConfigFile.py", ["add_core_level_Data"])
    data = {"Core levels": {}}
    for name in sheet_names:
        cfg["add_core_level_Data"](data, None, xlsx, name)
    out = {}
    finals = avantage.refresh_sheet_names(list(sheet_names))   # checked below
    for old, new in zip(sheet_names, finals):
        cl = dict(data["Core levels"][old])
        df = pd.read_excel(xlsx, sheet_name=old)
        cl["B.E."] = [be + 0 for be in df.iloc[:, 0].tolist()]
        cl["Raw Data"] = df.iloc[:, 1].tolist()
        out[new] = cl
    return out


def test_refresh_names_match_khervefitting_source():
    if not os.path.isdir(KF_LIB):
        pytest.skip("KherveFitting source not here")
    src = open(os.path.join(KF_LIB, "FileMenu", "Save.py"), encoding="utf-8").read()
    for snippet in ("if 'survey' in lower_name:", "elif 'wide' in lower_name:",
                    r"re.search(r'([A-Z][a-z]?)\s+(\d+[spdf])', old_name)",
                    r"re.search(r'([A-Z][a-z]?\d+[spdf])', new_name)",
                    "if old_name in wb.sheetnames and new_name not in wb.sheetnames:"):
        assert snippet in src
    assert avantage.refresh_name("XPS Survey") == "Survey"
    assert avantage.refresh_name("C1s2_1") == "C1s1"
    assert avantage.refresh_name("Cut-Off") == "Cut-Off"
    assert avantage.refresh_sheet_names(["XPS Survey", "Survey Scan", "O1s"]) == \
        ["Survey", "Survey Scan", "O1s"]


# ---------------------------------------------------------------------------
# Avantage
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("fname", sorted(AVANTAGE))
def test_avantage_real_files(fname):
    path = os.path.join(RAW, fname)
    if not os.path.exists(path):
        pytest.skip("sample not on this machine")
    fmts = vendors.formats_for(path)
    assert fmts and fmts[0].key == "avantage"
    res = avantage.read_avantage(path)
    assert [s["Name"] for s in res.spectra] == AVANTAGE[fname]
    assert {n for n, _r in res.dismissed} == {"Peak Table", "Titles", "Sheet1"}
    for sp in res.spectra:
        x, y = sp["B.E."], sp["Raw Data"]
        assert len(x) == len(y) > 50
        assert x[0] > x[-1]                         # BE high -> low, as exported
        assert min(y) > 0
        assert all(v == float(f"{v:.2f}") for v in x + y)
        assert sp["ExperimentalInfo"]["Source Gun Type"] == "Al K Alpha"
        assert "Sample Info" in sp["ExperimentalInfo"]
        assert "X_Label" not in sp


@need_kf
@pytest.mark.parametrize("fname", sorted(AVANTAGE))
def test_avantage_equals_khervefitting(fname, tmp_path, monkeypatch):
    path = os.path.join(RAW, fname)
    if not os.path.exists(path):
        pytest.skip("sample not on this machine")
    src = tmp_path / fname
    shutil.copy2(path, src)                       # KF writes <name>_Kfitting.xlsx beside it
    opened = []
    mrs = kf_functions("FileMenu/MRS_Import.py", ["extract_acquisition_parameters"])
    _stub_libraries(monkeypatch, **{
        "libraries.FileMenu.Open": {"open_xlsx_file": lambda w, p: opened.append(p)}})
    kf = kf_functions(
        "FileMenu/Avantage_Import.py",
        ["import_avantage_file_direct", "extract_title_sheet_info"],
        extract_acquisition_parameters=mrs["extract_acquisition_parameters"],
        extract_peak_fitting_from_peak_table=lambda wb: {},
        add_peak_fitting_data_after_load=lambda *a: None)
    monkeypatch.setattr("builtins.print", lambda *a, **k: None)
    kf["import_avantage_file_direct"](None, str(src))
    import openpyxl
    wb = openpyxl.load_workbook(opened[0], read_only=True)
    names = wb.sheetnames
    wb.close()
    ref = kf_read_back(opened[0], names)

    res = avantage.read_avantage(str(src))
    assert [s["Name"] for s in res.spectra] == list(ref)
    for sp in res.spectra:
        r = ref[sp["Name"]]
        assert sp["B.E."] == r["B.E."]
        assert sp["Raw Data"] == r["Raw Data"]
        assert sp["Corrected Data"] == r["Corrected Data"]
        assert sp["Transmission"] == r["Transmission"]
        assert sp["ExperimentalInfo"] == r["ExperimentalInfo"]


def test_avantage_sniff_rejects_khervefitting_workbook(tmp_path):
    if os.path.exists(KF_WORKBOOK):
        assert not avantage.sniff_avantage(KF_WORKBOOK)
        assert "avantage" not in [f.key for f in vendors.formats_for(KF_WORKBOOK)]
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "C1s"
    wb.active.append(["Binding Energy", "Raw Data"])
    p = str(tmp_path / "plain.xlsx")
    wb.save(p)
    assert "avantage" not in [f.key for f in vendors.formats_for(p)]


def _avantage_like(ws, be0, counts, n_samples=1):
    """A minimal Avantage spectrum sheet: parameters at H1.., data under 'eV'."""
    ws["H1"] = "Acquisition Parameters :"
    ws["H3"], ws["I3"] = "Parameter", "  "
    ws["H4"], ws["I4"] = "Number of Scans", 5
    ws["H5"], ws["I5"] = "Analyser Mode", "CAE : Pass Energy 20.0 eV"
    if n_samples > 1:
        ws["D8"] = n_samples
    ws["A15"] = "Binding Energy (E)"
    ws["A16"], ws["C16"] = "eV", "Counts / s"
    for i, c in enumerate(counts):
        ws.cell(row=17 + i, column=1, value=be0 - 0.1 * i)
        for k in range(n_samples):
            ws.cell(row=17 + i, column=3 + k, value=c * (k + 1) + 0.123456)


def test_avantage_synthetic_names_and_multisample(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    counts = [100 + 10 * np.sin(i / 3) for i in range(30)]
    _avantage_like(wb.create_sheet("XPS Survey"), 1200.0, counts)
    _avantage_like(wb.create_sheet("C1s Scan"), 295.0, counts)
    _avantage_like(wb.create_sheet("C1s Scan (2)"), 295.0, counts)
    _avantage_like(wb.create_sheet("Valence Band"), 10.0, counts)
    _avantage_like(wb.create_sheet("O1s Scan"), 540.0, counts, n_samples=2)
    _avantage_like(wb.create_sheet("C1s Scan B"), 295.0, counts)
    wb.create_sheet("Titles")["A1"] = "My sample"
    wb.create_sheet("Titles B")["A1"] = "Second sample"
    p = str(tmp_path / "synthetic.xlsx")
    wb.save(p)
    assert avantage.sniff_avantage(p)
    res = avantage.read_avantage(p)
    names = [s["Name"] for s in res.spectra]
    assert names == ["Survey", "C1s", "C1s2", "VB", "O1s", "O1s1", "C1s1"]
    by = {s["Name"]: s for s in res.spectra}
    # single-sample sheets: Raw Data kept to 2 decimals
    assert by["C1s"]["Raw Data"][0] == float(f"{counts[0] + 0.123456:.2f}")
    # multi-sample sheet: column B written f"{v}" -> reloaded unrounded
    assert by["O1s"]["Raw Data"][0] == counts[0] + 0.123456
    assert by["O1s1"]["Raw Data"][0] == counts[0] * 2 + 0.123456
    assert by["O1s"]["Corrected Data"][0] == float(f"{counts[0] + 0.123456:.2f}")
    assert by["C1s1"]["ExperimentalInfo"]["Sample Info"] == "Second sample"
    assert by["C1s"]["ExperimentalInfo"]["Number of Scans"] == "5"


@need_kf
def test_avantage_synthetic_equals_khervefitting(tmp_path, monkeypatch):
    test_avantage_synthetic_names_and_multisample(tmp_path)
    p = str(tmp_path / "synthetic.xlsx")
    opened = []
    mrs = kf_functions("FileMenu/MRS_Import.py", ["extract_acquisition_parameters"])
    _stub_libraries(monkeypatch, **{
        "libraries.FileMenu.Open": {"open_xlsx_file": lambda w, q: opened.append(q)}})
    kf = kf_functions(
        "FileMenu/Avantage_Import.py",
        ["import_avantage_file_direct", "extract_title_sheet_info"],
        extract_acquisition_parameters=mrs["extract_acquisition_parameters"],
        extract_peak_fitting_from_peak_table=lambda wb: {},
        add_peak_fitting_data_after_load=lambda *a: None)
    monkeypatch.setattr("builtins.print", lambda *a, **k: None)
    kf["import_avantage_file_direct"](None, p)
    import openpyxl
    wb = openpyxl.load_workbook(opened[0], read_only=True)
    names = wb.sheetnames
    wb.close()
    ref = kf_read_back(opened[0], names)
    res = avantage.read_avantage(p)
    assert [s["Name"] for s in res.spectra] == list(ref)
    for sp in res.spectra:
        assert sp["B.E."] == ref[sp["Name"]]["B.E."]
        assert sp["Raw Data"] == ref[sp["Name"]]["Raw Data"]
        assert sp["ExperimentalInfo"] == ref[sp["Name"]]["ExperimentalInfo"]


# ---------------------------------------------------------------------------
# VGD
# ---------------------------------------------------------------------------
@pytest.mark.skipif(not VGD_FILES, reason="no VGD samples on this machine")
@pytest.mark.parametrize("path", VGD_FILES, ids=lambda p: os.path.relpath(p, VGD_ROOT))
def test_vgd_real_files(path):
    fmts = vendors.formats_for(path)
    assert fmts and fmts[0].key == "vgd"
    res = vgd.read_vgd(path)
    assert len(res.spectra) == 1 and not res.dismissed
    sp = res.spectra[0]
    expected = vgd.extract_core_level_name(os.path.basename(path))
    assert sp["Name"] == expected
    assert re.match(r"^([A-Z][a-z]?\d[spdf]|XPS Survey)$", sp["Name"])
    x, y = sp["B.E."], sp["Raw Data"]
    assert len(x) == len(y) == int(sp["ExperimentalInfo"]["Number of Points"]) > 100
    assert x[0] > x[-1]
    assert min(y) > 0
    assert sp["ExperimentalInfo"]["Technique"] == "XPS"


@need_kf
@pytest.mark.skipif(not VGD_FILES, reason="no VGD samples on this machine")
@pytest.mark.parametrize("path", VGD_FILES, ids=lambda p: os.path.relpath(p, VGD_ROOT))
def test_vgd_equals_khervefitting(path, monkeypatch):
    kf = _kf_vgd(monkeypatch)
    parsed = kf["parse_vgd_file"](path)
    core = kf["extract_core_level_name"](os.path.basename(path))
    base = os.path.splitext(os.path.basename(path))[0]
    ref = dict(kf["build_vgd_core_levels"](parsed, core, base, []))
    res = vgd.read_vgd(path)
    assert [s["Name"] for s in res.spectra] == list(ref)
    for sp in res.spectra:
        r = ref[sp["Name"]]
        for key in ("B.E.", "Raw Data", "Corrected Data", "Transmission",
                    "ExperimentalInfo"):
            assert sp[key] == r[key], key


def _kf_vgd(monkeypatch):
    import olefile
    cfg = kf_functions("ConfigFile.py", ["build_core_level_Data"])
    _stub_libraries(monkeypatch, **{
        "libraries.ConfigFile": {"build_core_level_Data": cfg["build_core_level_Data"]}})
    return kf_functions(
        "FileMenu/VGD_Import.py",
        ["parse_vgd_file", "calculate_vgd_data", "extract_core_level_name",
         "vgd_spectrum_metadata", "vgd_map_pixel_data", "_unique_name",
         "build_vgd_map_metadata", "build_vgd_average_metadata",
         "build_vgd_spectrum_core_level", "build_vgd_xy_map_core_levels",
         "build_vgd_core_levels"], olefile=olefile)


def _parsed(num_spectra, n_ke, area=None):
    rng = np.random.default_rng(1)
    if area:
        n_x, n_y = area
        inten = [list(1000 + 100 * rng.random(n_ke)) for _ in range(n_x * n_y)]
    elif num_spectra > 1:
        inten = [list(1000 + 100 * rng.random(n_ke)) for _ in range(num_spectra)]
    else:
        inten = list(1000 + 100 * rng.random(n_ke))
    return {
        'intensities': inten, 'ke_start': 1190.0, 'ke_step': 0.1, 'num_points': n_ke,
        'total_points': n_ke * max(1, num_spectra), 'num_spectra': num_spectra,
        'is_area_scan': bool(area), 'n_x': area[0] if area else None,
        'n_y': area[1] if area else None, 'x_start': 0.0, 'x_step': 50.0 if area else None,
        'y_start': 0.0, 'y_step': 50.0 if area else None,
        'source_energy': 1486.6800537109375, 'txf_coeffs': [], 'pass_energy': 20.0,
        'work_fn': 4.5, 'dwell_time': 0.05, 'periods': 10,
        'metadata': {'title': 'C1s Scan', 'subject': 'S', 'author': 'a',
                     'create_time': '2026-01-01 10:00:00', 'saved_time': ''}}


@need_kf
@pytest.mark.parametrize("kind", ["multi", "area"])
def test_vgd_multi_and_area_equal_khervefitting(kind, monkeypatch):
    kf = _kf_vgd(monkeypatch)
    parsed = _parsed(3, 40) if kind == "multi" else _parsed(4, 40, area=(2, 2))
    ref = dict(kf["build_vgd_core_levels"](parsed, "C1s", "C1s Scan", []))
    spectra, dismissed = vgd.vgd_core_levels(parsed, "C1s", "C1s Scan", [])
    if kind == "multi":
        assert [s["Name"] for s in spectra] == ["C1s", "C1s1", "C1s2"] == list(ref)
        assert not dismissed
    else:
        assert list(ref) == ["C1s~Map", "C1s"]
        assert [s["Name"] for s in spectra] == ["C1s"]
        assert [n for n, _r in dismissed] == ["C1s~Map"]
    for sp in spectra:
        r = ref[sp["Name"]]
        for key in ("B.E.", "Raw Data", "Corrected Data", "Transmission",
                    "ExperimentalInfo"):
            assert sp[key] == r[key], key


# ---------------------------------------------------------------------------
# AVG (synthetic — no sample exists)
# ---------------------------------------------------------------------------
def _write_avg(path, spectra, n=60, ke0=1190.0, step=0.1, etch=None):
    lines = [
        "$PROPERTIES=",
        "DS_EXT_SUPROPID_TITLE : Title : VT_BSTR = 'C1s Snap'",
        "DS_EXT_SUPROPID_SUBJECT : Subject : VT_BSTR = 'Sample-7'",
        "DS_EXT_SUPROPID_AUTHOR : Author : VT_BSTR = 'engineer'",
        "DS_EXT_SUPROPID_CREATED : Created : VT_DATE = 16/09/2026 10:11:12",
        "DS_GEPROPID_INSTRUMENT : Instrument : VT_BSTR = 'K-Alpha'",
        "DS_SOPROPID_ENERGY : Source energy : VT_R4 = 1486.68",
        "DS_ANPROPID_PASS : Pass energy : VT_R4 = 20",
        "DS_ANPROPID_WORK_FTN : Work function : VT_R4 = 4.5",
        "DS_ACPROPID_ACQ_TIME : Dwell : VT_R4 = 0.05",
        "DS_ACPROPID_PERIODS : Periods : VT_I4 = 7",
        "DS_ANPROPID_LENS_MODE_NAME : Lens : VT_BSTR = 'Standard'",
        "DS_ANPROPID_TXFN_COEFF[0] : TXF : VT_R4 = 4.25",
        "$DATAAXES=%d,#empty#" % (2 if etch else 1),
        "$SPACEAXES=%d" % (2 if etch else 1),
        f"   0=   {ke0},   {step},   {n},   ENERGY,   LINEAR,   'KE',   'eV',   'Kinetic Energy'",
    ]
    if etch:
        lines.append(f"   1=   0.0,   30.0,   {len(spectra)},   ETCHTIME,   LINEAR,   "
                     "'t',   's',   'Etch Time'")
        for i, v in enumerate(etch):
            lines.append(f"$AXISVALUE= DATAXIS=1 SPACEAXIS=1 LABEL='Etch Time' "
                         f"POINT={i} VALUE={v:.6f};")
    for counts in spectra:
        lines.append("$DATA=*")
        for i in range(0, len(counts), 4):
            chunk = ",  ".join(f"{c:.6f}" for c in counts[i:i + 4])
            lines.append(f"LIST@   {i}=  {chunk}")
        lines.append(";")
    with open(path, "w", encoding="latin-1") as fh:
        fh.write("\n".join(lines) + "\n")


def _counts(n, k=1.0):
    return [k * (500 + 300 * np.exp(-((i - n / 2) / 6.0) ** 2)) + 0.3141 * i
            for i in range(n)]


def test_avg_single_and_series(tmp_path):
    p = str(tmp_path / "C1s Snap.avg")
    _write_avg(p, [_counts(60)])
    fmts = vendors.formats_for(p)
    assert fmts and fmts[0].key == "avg"
    res = avg.read_avg(p)
    assert [s["Name"] for s in res.spectra] == ["C1s"]
    sp = res.spectra[0]
    assert len(sp["B.E."]) == 60 and sp["B.E."][0] > sp["B.E."][-1]
    assert sp["B.E."][0] == float(f"{1486.68 - 1190.0:.2f}")
    assert sp["Raw Data"][5] == float(f"{_counts(60)[5] / (0.05 * 7):.2f}")
    assert sp["Corrected Data"][5] == float(f"{_counts(60)[5]:.2f}")
    assert sp["Transmission"][0] == 0.35
    info = sp["ExperimentalInfo"]
    assert info["Sample ID"] == "Sample-7" and info["Periods"] == "7"
    assert info["TXF Coefficients"] == "4.250000"

    p2 = str(tmp_path / "XPS Survey.avg")
    _write_avg(p2, [_counts(60), _counts(60, 2), _counts(60, 3)], etch=[0, 30, 60])
    res = avg.read_avg(p2)
    assert [s["Name"] for s in res.spectra] == ["Survey", "Survey1", "Survey2"]
    assert res.spectra[2]["ExperimentalInfo"]["Etch Time (s)"] == "60.0000"
    assert min(res.spectra[1]["Raw Data"]) > 0


@need_kf
def test_avg_equals_khervefitting(tmp_path):
    import openpyxl
    p = str(tmp_path / "Ti2p Snap.avg")
    _write_avg(p, [_counts(57), _counts(57, 1.5)], n=57, etch=[0, 30])
    kf = kf_functions("FileMenu/AVG_Import.py",
                      ["parse_avg_file", "calculate_avg_data", "extract_core_level_name",
                       "_get_bstr", "_get_r4", "_get_i4", "_get_date", "_write_sheet"])
    parsed = kf["parse_avg_file"](p)
    core = kf["extract_core_level_name"](os.path.basename(p))
    # open_avg_file: one sheet per spectrum, then _load_into_kherve
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = []
    for idx in range(parsed["num_spectra"]):
        calc = kf["calculate_avg_data"](parsed, idx)
        name = core if idx == 0 else f"{core}{idx}"
        names.append(name)
        kf["_write_sheet"](wb.create_sheet(name), core, calc, parsed, idx, "Ti2p Snap", name)
    xlsx = str(tmp_path / "Ti2p Snap.xlsx")
    wb.save(xlsx)
    ref = kf_read_back(xlsx, names)

    res = avg.read_avg(p)
    assert [s["Name"] for s in res.spectra] == list(ref) == ["Ti2p", "Ti2p1"]
    for sp in res.spectra:
        r = ref[sp["Name"]]
        for key in ("B.E.", "Raw Data", "Corrected Data", "Transmission",
                    "ExperimentalInfo"):
            assert sp[key] == r[key], key


# ---------------------------------------------------------------------------
# Avantage .xls route — no .xls writer is installed, so xlrd.open_workbook is
# replaced by a book built from an openpyxl workbook (same xlrd API subset).
# ---------------------------------------------------------------------------
class _FakeXlsSheet:
    def __init__(self, ws):
        self.rows = [[("" if v is None else (float(v) if isinstance(v, (int, float)) else v))
                      for v in r] for r in ws.iter_rows(values_only=True)]
        self.nrows = len(self.rows)
        self.ncols = max((len(r) for r in self.rows), default=0)

    def cell_value(self, r, c):
        return self.rows[r][c]

    def row_values(self, r):
        return list(self.rows[r])


class _FakeXlsBook:
    def __init__(self, wb):
        self._sheets = {ws.title: _FakeXlsSheet(ws) for ws in wb.worksheets}

    def sheet_names(self):
        return list(self._sheets)

    def sheet_by_name(self, name):
        return self._sheets[name]

    def release_resources(self):
        pass


def _fake_xls(tmp_path, monkeypatch):
    import openpyxl
    import xlrd
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    counts = [200 + 20 * np.cos(i / 4) for i in range(25)]
    _avantage_like(wb.create_sheet("Survey"), 1200.0, counts)
    _avantage_like(wb.create_sheet("Survey (2)"), 1200.0, counts)
    _avantage_like(wb.create_sheet("Fe2p Scan"), 730.0, counts)
    _avantage_like(wb.create_sheet("Fermi Edge"), 3.0, counts)
    _avantage_like(wb.create_sheet("Fe2p Scan A"), 730.0, counts)
    wb.create_sheet("Titles")["A1"] = "Xls sample"
    book = _FakeXlsBook(wb)
    monkeypatch.setattr(xlrd, "open_workbook", lambda *a, **k: book)
    p = tmp_path / "legacy.xls"
    p.write_bytes(b"not read: xlrd is faked")
    return str(p), counts


def test_avantage_xls_route(tmp_path, monkeypatch):
    p, counts = _fake_xls(tmp_path, monkeypatch)
    assert [f.key for f in vendors.formats_for(p)][:1] == ["avantage"]
    res = avantage.read_avantage(p)
    assert [s["Name"] for s in res.spectra] == ["Survey", "Survey2", "Fe2p", "Fermi", "Fe2p1"]
    sp = res.spectra[2]
    assert sp["Raw Data"][3] == float(f"{counts[3] + 0.123456:.2f}")
    assert sp["ExperimentalInfo"]["Sample Info"] == "Xls sample"
    assert sp["ExperimentalInfo"]["Number of Scans"] == "5.0"     # xlrd floats


@need_kf
def test_avantage_xls_equals_khervefitting(tmp_path, monkeypatch):
    p, _counts = _fake_xls(tmp_path, monkeypatch)
    opened = []
    mrs = kf_functions("FileMenu/MRS_Import.py", ["extract_acquisition_parameters"])
    _stub_libraries(monkeypatch, **{
        "libraries.FileMenu.Open": {"open_xlsx_file": lambda w, q: opened.append(q)}})
    kf = kf_functions(
        "FileMenu/Avantage_Import.py",
        ["import_avantage_file_direct_xls", "extract_title_sheet_info_xls"],
        extract_acquisition_parameters=mrs["extract_acquisition_parameters"])
    monkeypatch.setattr("builtins.print", lambda *a, **k: None)
    kf["import_avantage_file_direct_xls"](None, p)
    import openpyxl
    wb = openpyxl.load_workbook(opened[0], read_only=True)
    names = wb.sheetnames
    wb.close()
    ref = kf_read_back(opened[0], names)
    res = avantage.read_avantage(p)
    assert [s["Name"] for s in res.spectra] == list(ref)
    for sp in res.spectra:
        r = ref[sp["Name"]]
        for key in ("B.E.", "Raw Data", "Corrected Data", "Transmission",
                    "ExperimentalInfo"):
            assert sp[key] == r[key], key
