"""Thermo Avantage Excel exports (.xlsx / .xls with a "Titles" sheet).

Reproduces KherveFitting's ``Avantage_Import.import_avantage_file_direct``
(.xlsx) and ``import_avantage_file_direct_xls`` (.xls), with
``extract_title_sheet_info`` / ``extract_title_sheet_info_xls`` and
``MRS_Import.extract_acquisition_parameters`` for the metadata.

Route reproduced: the **workbook route** (it is the only one Avantage has).
KherveFitting copies every spectrum sheet into ``<file>_Kfitting.xlsx``
(column A = BE as ``f"{v:.2f}"`` text, column B = the counts of column C),
opens it with ``Open.open_xlsx_file`` -> ``ConfigFile.add_core_level_Data``
and then ``Save.refresh_sheets``, which

* renames sheets with its normaliser (``"C1s2_1"`` -> ``"C1s1"`` ...; a
  rename is skipped when the target name is already taken), and
* reloads ``'B.E.'`` / ``'Raw Data'`` from columns A / B as stored
  (text parsed by pandas — reproduced with ``pd.to_numeric``, which can
  differ from ``float()`` in the last bit).

So here: ``'B.E.'`` = BE to 2 decimals; ``'Raw Data'`` = counts to 2
decimals (single-sample sheets are written ``.2f``) or unrounded
(multi-sample sheets are written ``f"{v}"``); ``'Corrected Data'`` = counts
to 2 decimals (add_core_level_Data falls back to column B);
``'Transmission'`` = 1.  ``ExperimentalInfo`` = the Titles-sheet lines
followed by the sheet's "Acquisition Parameters" block, as KherveFitting
writes them into the "Experimental Description" column.

Not ported: the "Peak Table" fit extraction (``extract_peak_fitting_*``,
``add_peak_fitting_data_after_load``) — KherveNoise only needs the spectra.
Deliberate deviation: in the .xls route KherveFitting looks a duplicate
sheet up by its requested name after openpyxl renamed it, so the duplicate's
data is appended to the first sheet and the renamed one stays empty; here
each duplicate keeps its own data under openpyxl's name (``C1s1`` ...).
Rows KherveFitting would leave as blank gaps are dropped, as
add_core_level_Data does.

Copyright (C) 2026 Gwilherm Kerherve
Licensed under the GNU General Public License v3.0 (see LICENSE).
"""

import os
import re

from . import Format
from ..document import make_spectrum
from ..importers import ImportResult, _r2

#: Sheet-name terms that make a sheet a spectrum (``import_avantage_file_direct``).
_TERMS_XLSX = ['valence', 'valence band', 'valence scan', 'vb scan', 'vb',
               'valence band scan', 'val', 'fermi', 'fermi scan',
               'cut-off', 'cutoff', 'cut off']
#: ``import_avantage_file_direct_xls`` has the same list without 'val'.
_TERMS_XLS = ['valence', 'valence band', 'valence scan', 'vb scan', 'vb',
              'valence band scan', 'fermi', 'fermi scan',
              'cut-off', 'cutoff', 'cut off']
_VB_TERMS = ['valence', 'valence band', 'valence scan', 'vb scan', 'vb',
             'valence band scan']
_FERMI_TERMS = ['fermi', 'fermi scan']
_CUTOFF_TERMS = ['cut-off', 'cutoff', 'cut off']


# ---------------------------------------------------------------------------
# Sniff
# ---------------------------------------------------------------------------
def _is_titles(name):
    return name == "Titles" or bool(re.match(r'^Titles [A-Za-z]$', name))


def sheet_names(path):
    if str(path).lower().endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def sniff_avantage(path):
    """True only for a workbook that has Avantage's "Titles" sheet."""
    return any(_is_titles(n) for n in sheet_names(path))


# ---------------------------------------------------------------------------
# Sheet-name handling shared by the workbook routes
# ---------------------------------------------------------------------------
def refresh_name(old_name):
    """``Save.refresh_sheets``' sheet-name normaliser, verbatim."""
    new_name = old_name
    lower_name = old_name.lower()
    if 'survey' in lower_name:
        new_name = 'Survey'
    elif 'xps survey' in lower_name:
        new_name = 'Survey'
    elif 'survey scan' in lower_name:
        new_name = 'Survey'
    elif 'wide' in lower_name:
        new_name = 'Wide'
    elif 'wide scan' in lower_name:
        new_name = 'Wide'
    else:
        match = re.search(r'([A-Z][a-z]?)\s+(\d+[spdf])', old_name)
        if match:
            element, orbital = match.groups()
            new_name = f"{element}{orbital}"
        match = re.search(r'([A-Z][a-z]?\d+[spdf])', new_name)
        if match and len(new_name) > len(match.group(1)):
            new_name = match.group(1)
    suffix_match = re.search(r'(\d+)$', old_name)
    if suffix_match and not re.search(r'\d+$', new_name):
        new_name = f"{new_name}{suffix_match.group(1)}"
    return new_name


def refresh_sheet_names(names):
    """Final names after ``refresh_sheets``: each sheet is renamed in turn
    unless its new name is already a sheet of the workbook."""
    current = list(names)
    for i, old in enumerate(names):
        new = refresh_name(old)
        if new != old and new not in current:
            current[i] = new
    return current


def _create_sheet_name(taken, name):
    """The title openpyxl's ``create_sheet`` actually gives (it appends a
    number to a name already in the workbook, case-insensitively)."""
    from openpyxl.workbook.child import avoid_duplicate_name
    return avoid_duplicate_name(taken, name)


# ---------------------------------------------------------------------------
# Metadata
# ---------------------------------------------------------------------------
def extract_acquisition_parameters(cell, max_row, max_column):
    """``MRS_Import.extract_acquisition_parameters`` over ``cell(row, col)``
    (1-based, returns the value or None)."""
    parameters = {}
    acquisition_col = None
    for row in range(1, min(5, max_row + 1)):
        for col in range(1, max_column + 1):
            v = cell(row, col)
            if v and "Acquisition Parameters" in str(v):
                acquisition_col = col
                break
        if acquisition_col:
            break
    if acquisition_col is None:
        for row in range(1, min(5, max_row + 1)):
            for col in range(1, max_column + 1):
                v = cell(row, col)
                if v and "Parameters" in str(v):
                    acquisition_col = col
                    break
            if acquisition_col:
                break
    if acquisition_col is None:
        return parameters

    for row in range(3, min(15, max_row + 1)):
        param_name = cell(row, acquisition_col)
        param_value = cell(row, acquisition_col + 1)
        param_value_j = cell(row, acquisition_col + 2)
        if param_name is None:
            continue
        param_name_str = str(param_name).strip()
        if not param_name_str or param_name_str in ['Parameter', 'Parameters', '']:
            continue
        value_parts = []
        if param_value is not None and str(param_value).strip():
            value_parts.append(str(param_value).strip())
        if param_value_j is not None and str(param_value_j).strip():
            value_parts.append(str(param_value_j).strip())
        if value_parts:
            parameters[param_name_str] = " ".join(value_parts)
    return parameters


def _title_sheet_name(sample_suffix):
    return "Titles" if sample_suffix is None else f"Titles {sample_suffix}"


def extract_title_sheet_info(wb, sample_suffix=None):
    """``extract_title_sheet_info`` (openpyxl): A1, A3, A5 of "Titles [X]"."""
    title_info = {}
    name = _title_sheet_name(sample_suffix)
    if name in wb.sheetnames:
        ws = wb[name]
        a1 = ws.cell(row=1, column=1).value
        a3 = ws.cell(row=3, column=1).value
        a5 = ws.cell(row=5, column=1).value
        if a1:
            title_info['Sample Info'] = str(a1).strip()
        if a3:
            title_info['Additional Info'] = str(a3).strip()
        if a5:
            title_info['Notes'] = str(a5).strip()
    return title_info


def extract_title_sheet_info_xls(wb_xls, sample_suffix=None):
    """``extract_title_sheet_info_xls`` (xlrd)."""
    title_info = {}
    name = _title_sheet_name(sample_suffix)
    if name in wb_xls.sheet_names():
        ts = wb_xls.sheet_by_name(name)
        for (r, key) in ((0, 'Sample Info'), (2, 'Additional Info'), (4, 'Notes')):
            try:
                v = ts.cell_value(r, 0)
                if v:
                    title_info[key] = str(v).strip()
            except Exception:  # noqa: BLE001 — KherveFitting's bare except
                pass
    return title_info


def _exp_info(title_info, acquisition_params):
    """What add_core_level_Data reads back from the description column."""
    info = {}
    for k, v in list(title_info.items()) + list(acquisition_params.items()):
        info[str(k).strip()] = '' if v is None else str(v).strip()
    return info


# ---------------------------------------------------------------------------
# Grouping (shared)
# ---------------------------------------------------------------------------
def _sample_groups(names, terms):
    groups = {}
    for sheet_name in names:
        if "Survey" in sheet_name or "Scan" in sheet_name or \
                any(term in sheet_name.lower() for term in terms):
            parts = sheet_name.split()
            if len(parts) >= 2 and len(parts[-1]) == 1 and parts[-1].isalpha():
                suffix = parts[-1]
            else:
                suffix = None
            groups.setdefault(suffix, []).append(sheet_name)
    order = sorted(groups.keys(), key=lambda x: (x is not None, x))
    return groups, order


def _not_spectrum_reason(name):
    if _is_titles(name):
        return "Avantage title sheet (kept as sample information)"
    if name.lower() == 'peak table':
        return "Avantage peak table (fit results, not a spectrum)"
    return "not an Avantage spectrum sheet"


def _as_float(v):
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


# ---------------------------------------------------------------------------
# .xlsx  (import_avantage_file_direct)
# ---------------------------------------------------------------------------
def _entries_xlsx(path, dismissed):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    try:
        groups, order = _sample_groups(wb.sheetnames, _TERMS_XLSX)
        grouped = {n for g in groups.values() for n in g}
        for n in wb.sheetnames:
            if n not in grouped:
                dismissed.append((n, _not_spectrum_reason(n)))

        entries, created = [], []

        def new_sheet(name):
            title = _create_sheet_name(created, name)
            created.append(title)
            return title

        sample_counter = 0
        for sample_suffix in order:
            title_info = extract_title_sheet_info(wb, sample_suffix)
            vb_count = fermi_count = cutoff_count = 0
            for sheet_name in groups[sample_suffix]:
                sheet = wb[sheet_name]
                acq = extract_acquisition_parameters(
                    lambda r, c, s=sheet: s.cell(row=r, column=c).value,
                    sheet.max_row, sheet.max_column)
                lower_name = sheet_name.lower()
                if "Survey" in sheet_name or "survey" in sheet_name:
                    base_name = "Survey"
                elif any(t in lower_name for t in _VB_TERMS):
                    base_name = "VB" if vb_count == 0 else f"VB{vb_count + 1}"
                    vb_count += 1
                elif any(t in lower_name for t in _FERMI_TERMS):
                    base_name = "Fermi" if fermi_count == 0 else f"Fermi{fermi_count + 1}"
                    fermi_count += 1
                elif any(t in lower_name for t in _CUTOFF_TERMS):
                    base_name = "Cut-Off" if cutoff_count == 0 else f"Cut-Off{cutoff_count + 1}"
                    cutoff_count += 1
                else:
                    base_name = sheet_name.split()[0]
                final_name = base_name if sample_counter == 0 else f"{base_name}{sample_counter}"

                multi_sample, num_samples = False, 1
                d8 = sheet.cell(row=8, column=4).value
                if d8 is not None:
                    try:
                        num_samples = int(d8)
                        multi_sample = num_samples > 1
                    except (ValueError, TypeError):
                        pass

                start_row = 19
                for row_idx in range(1, sheet.max_row + 1):
                    if sheet.cell(row=row_idx, column=1).value == "eV":
                        start_row = row_idx + 1
                        break

                info = (_exp_info(title_info, acq) if (acq or title_info) else None)
                if multi_sample:
                    for sample_idx in range(num_samples):
                        if sample_counter == 0:
                            sample_name = f"{base_name}{sample_idx if sample_idx > 0 else ''}"
                        elif sample_idx == 0:
                            sample_name = f"{base_name}{sample_counter}"
                        else:
                            sample_name = f"{base_name}{sample_counter}_{sample_idx}"
                        xs, ys = [], []
                        for row in range(start_row, sheet.max_row + 1):
                            be = sheet.cell(row=row, column=1).value
                            it = sheet.cell(row=row, column=3 + sample_idx).value
                            if _as_float(be) is None or _as_float(it) is None:
                                continue
                            xs.append(float(be))
                            ys.append(float(it))
                        # Column B written as f"{v}": refresh_sheets reloads it unrounded.
                        entries.append((new_sheet(sample_name), xs, ys, False, info, sheet_name))
                else:
                    new_name = final_name
                    number_match = re.search(r'\((\d+)\)', sheet_name)
                    if number_match:
                        number = number_match.group(1)
                        new_name = (f"{base_name}{number}" if sample_counter == 0
                                    else f"{base_name}{sample_counter}_{number}")
                    xs, ys = [], []
                    for row in range(start_row, sheet.max_row + 1):
                        be = sheet.cell(row=row, column=1).value
                        it = sheet.cell(row=row, column=3).value
                        if _as_float(be) is None or _as_float(it) is None:
                            continue
                        xs.append(float(be))
                        ys.append(float(it))
                    entries.append((new_sheet(new_name), xs, ys, True, info, sheet_name))
            sample_counter += 1
        return entries
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# .xls  (import_avantage_file_direct_xls)
# ---------------------------------------------------------------------------
class _Grid:
    """The temporary openpyxl sheet KherveFitting fills from the first
    25 x 20 non-blank cells to run extract_acquisition_parameters on."""

    def __init__(self, sheet):
        self.cells = {}
        for row in range(min(25, sheet.nrows)):
            for col in range(min(20, sheet.ncols)):
                v = sheet.cell_value(row, col)
                if v is not None and str(v).strip():
                    self.cells[(row + 1, col + 1)] = v
        self.max_row = max((r for r, _c in self.cells), default=1)
        self.max_column = max((c for _r, c in self.cells), default=1)

    def __call__(self, row, col):
        return self.cells.get((row, col))


def _entries_xls(path, dismissed):
    import xlrd
    wb_xls = xlrd.open_workbook(path)
    groups, order = _sample_groups(wb_xls.sheet_names(), _TERMS_XLS)
    grouped = {n for g in groups.values() for n in g}
    for n in wb_xls.sheet_names():
        if n not in grouped:
            dismissed.append((n, _not_spectrum_reason(n)))

    entries, created = [], []
    sample_counter = 0
    for sample_suffix in order:
        title_info = extract_title_sheet_info_xls(wb_xls, sample_suffix)
        vb_count = fermi_count = cutoff_count = 0
        for sheet_name in groups[sample_suffix]:
            sheet = wb_xls.sheet_by_name(sheet_name)
            grid = _Grid(sheet)
            acq = extract_acquisition_parameters(grid, grid.max_row, grid.max_column)
            lower_name = sheet_name.lower()
            if "Survey" in sheet_name or "survey" in sheet_name:
                number_match = re.search(r'\((\d+)\)', sheet_name)
                if number_match:
                    number = number_match.group(1)
                    new_name = (f"Survey{number}" if sample_counter == 0
                                else f"Survey{sample_counter}_{number}")
                else:
                    new_name = "Survey" if sample_counter == 0 else f"Survey{sample_counter}"
            elif any(t in lower_name for t in _VB_TERMS):
                base_name = "VB" if vb_count == 0 else f"VB{vb_count + 1}"
                vb_count += 1
                new_name = base_name if sample_counter == 0 else f"{base_name}{sample_counter}"
            elif any(t in lower_name for t in _FERMI_TERMS):
                base_name = "Fermi" if fermi_count == 0 else f"Fermi{fermi_count + 1}"
                fermi_count += 1
                new_name = base_name if sample_counter == 0 else f"{base_name}{sample_counter}"
            elif any(t in lower_name for t in _CUTOFF_TERMS):
                base_name = "Cut-Off" if cutoff_count == 0 else f"Cut-Off{cutoff_count + 1}"
                cutoff_count += 1
                new_name = base_name if sample_counter == 0 else f"{base_name}{sample_counter}"
            else:
                element = sheet_name.split()[0]
                number_match = re.search(r'\((\d+)\)', sheet_name)
                if number_match:
                    number = number_match.group(1)
                    new_name = (f"{element}{number}" if sample_counter == 0
                                else f"{element}{sample_counter}_{number}")
                else:
                    new_name = element if sample_counter == 0 else f"{element}{sample_counter}"

            start_row = 17
            for row_idx in range(sheet.nrows):
                if sheet.cell_value(row_idx, 0) == "eV":
                    start_row = row_idx + 1
                    break
            xs, ys = [], []
            for row_idx in range(start_row, sheet.nrows):
                row_values = sheet.row_values(row_idx)
                if len(row_values) < 3:
                    continue
                be, it = row_values[0], row_values[2]
                if _as_float(be) is None or _as_float(it) is None:
                    continue
                if be == xlrd.empty_cell.value or it == xlrd.empty_cell.value:
                    continue
                xs.append(float(be))
                ys.append(float(it))
            # The .xls route only writes the description when there are
            # acquisition parameters.
            info = _exp_info(title_info, acq) if acq else None
            title = _create_sheet_name(created, new_name)
            created.append(title)
            entries.append((title, xs, ys, True, info, sheet_name))
        sample_counter += 1
    return entries


def _reloaded(texts):
    """Columns A / B hold text; refresh_sheets reloads them through pandas,
    whose number parser can differ from float() in the last bit."""
    import pandas as pd
    return [float(v) for v in pd.to_numeric(pd.Series(texts, dtype=object)).tolist()]


# ---------------------------------------------------------------------------
def read_avantage(path):
    """Spectra of an Avantage workbook, as KherveFitting ends up with them."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"The file {path} does not exist.")
    result = ImportResult(source=path)
    if str(path).lower().endswith('.xls'):
        entries = _entries_xls(path, result.dismissed)
    else:
        entries = _entries_xlsx(path, result.dismissed)

    finals = refresh_sheet_names([e[0] for e in entries])
    for final, (_title, xs, ys, y_rounded, info, src) in zip(finals, entries):
        if len(xs) < 2:
            result.dismissed.append((src, "no numeric data"))
            continue
        x = _reloaded([f"{v:.2f}" for v in xs])
        y = _reloaded([f"{v:.2f}" if y_rounded else f"{v}" for v in ys])
        sp = make_spectrum(final, x, y,
                           raw=[_r2(v) for v in ys], transmission=[1.0] * len(ys),
                           info=info)
        result.spectra.append(sp)
    return result


FORMATS = [
    Format(key="avantage", label="Avantage data file (.xlsx, .xls)", group="Thermo",
           extensions=('.xlsx', '.xls'), reader=read_avantage,
           sniff=sniff_avantage, priority=10),
]
