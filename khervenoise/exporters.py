"""Write spectra out — a KherveFitting workbook, or plain text.  Qt-free.

The workbook uses KherveFitting's own sheet layout, so it opens there
directly (File ▸ Open): one sheet per spectrum, row 1 =
``BE | Corrected Data | Raw Data | Transmission``.  As in KherveFitting,
column B is the trace that gets plotted and fitted (our 'Raw Data' — the
denoised values on a denoised copy) and column C the uncorrected counts.
Spectrum metadata goes in the "Experimental Description" block at column
50, where KherveFitting looks for it.

Copyright (C) 2026 Gwilherm Kerherve
Licensed under the GNU General Public License v3.0 (see LICENSE).
"""

import csv
import os
import re

from .document import axis_labels, is_xps_like

_EXP_COL = 50
_BAD_SHEET_CHARS = re.compile(r'[\[\]\*\?/\\:]')


def _plain(label):
    text = str(label or '')
    text = re.sub(r'[\^_]\{([^}]*)\}', r'\1', text)
    text = re.sub(r'[\^_](\w)', r'\1', text)
    text = re.sub(r'\\([A-Za-z]+)', r'\1', text)
    return text.replace('$', '').strip()


def sheet_title(name, taken):
    """An Excel-legal, unique sheet name (31 chars, no []*?/\\:)."""
    base = _BAD_SHEET_CHARS.sub('_', str(name))[:31] or "Sheet"
    title, i = base, 1
    while title.lower() in taken:
        suffix = str(i)
        title = base[:31 - len(suffix)] + suffix
        i += 1
    taken.add(title.lower())
    return title


def _x_header(sp):
    name = sp.get('Name', '')
    if is_xps_like(name, sp) and not sp.get('X_Label'):
        return "BE"
    return _plain(axis_labels(name, sp)[0])


def _exp_rows(sp):
    rows = dict(sp.get('ExperimentalInfo') or {})
    den = sp.get('Denoise')
    if den:
        rows['Denoised from'] = den.get('source', '')
        for k, v in (den.get('params') or {}).items():
            rows[f'Denoise {k}'] = v
    return rows


def write_khervefitting_xlsx(path, spectra):
    """Every spectrum in *spectra* (list of dicts) as a KherveFitting sheet.
    Returns the sheet names used."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    taken, titles = set(), []
    for sp in spectra:
        title = sheet_title(sp['Name'], taken)
        titles.append(title)
        ws = wb.create_sheet(title)
        ws.append([_x_header(sp), "Corrected Data", "Raw Data", "Transmission"])
        xs = sp['B.E.']
        ys = sp['Raw Data']
        cs = sp.get('Corrected Data') or ys
        ts = sp.get('Transmission') or [1.0] * len(xs)
        for x, y, c, t in zip(xs, ys, cs, ts):
            ws.append([x, y, c, t])
        rows = _exp_rows(sp)
        if rows:
            ws.cell(row=1, column=_EXP_COL, value="Experimental Description")
            for j, (k, v) in enumerate(rows.items()):
                ws.cell(row=j + 2, column=_EXP_COL, value=str(k))
                ws.cell(row=j + 2, column=_EXP_COL + 1,
                        value=v if isinstance(v, (int, float)) else str(v))
    if not titles:
        raise ValueError("Nothing to export.")
    wb.save(path)
    return titles


def write_text(path, spectra, delimiter=None):
    """Columns side by side: x, y for each spectrum (CSV when the path
    ends in .csv, tab-separated otherwise)."""
    if delimiter is None:
        delimiter = ',' if path.lower().endswith('.csv') else '\t'
    if not spectra:
        raise ValueError("Nothing to export.")
    header, cols = [], []
    for sp in spectra:
        header += [f"{sp['Name']} {_x_header(sp)}", f"{sp['Name']} {_plain(axis_labels(sp['Name'], sp)[1])}"]
        cols += [sp['B.E.'], sp['Raw Data']]
    n = max(len(c) for c in cols)
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh, delimiter=delimiter)
        w.writerow(header)
        for i in range(n):
            w.writerow([c[i] if i < len(c) else '' for c in cols])
    return os.path.basename(path)
